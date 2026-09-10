import openpyxl,re
from pathlib import Path
p=Path(__file__).resolve().parent
a=openpyxl.load_workbook(p/'Ticket_Labour_Invoice_月度核对.xlsx',data_only=True)
b=openpyxl.load_workbook(p/'Ticket_Labour_Invoice_Monthly_Reconciliation_EN.xlsx',data_only=True)
checked=0
for sa,sb in zip(a,b):
    assert (sa.max_row,sa.max_column)==(sb.max_row,sb.max_column)
    for row in sa:
        for ca in row:
            cb=sb[ca.coordinate]
            if isinstance(ca.value,(float,int)):
                assert isinstance(cb.value,(float,int)) and abs(ca.value-cb.value)<1e-8,(sb.title,cb.coordinate,ca.value,cb.value)
                checked+=1
            assert not(isinstance(cb.value,str) and re.search(r'[\u4e00-\u9fff]',cb.value)),(sb.title,cb.coordinate,cb.value)
            assert cb.data_type!='e',(sb.title,cb.coordinate,cb.value)
            if ca.is_date:assert ca.value==cb.value
f=openpyxl.load_workbook(p/'Ticket_Labour_Invoice_Monthly_Reconciliation_EN.xlsx',data_only=False)
assert sum(len(s.tables) for s in f)==3
assert f['Ticket Details']['G11'].value=='0090040857'
print('Verified English labels, dates, tables and invoice identifiers. Numeric cells unchanged:',checked)
print('Sheets:',b.sheetnames)
