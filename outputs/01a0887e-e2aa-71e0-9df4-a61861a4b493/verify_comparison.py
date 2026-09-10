import openpyxl,json,re
from pathlib import Path
p=Path(__file__).resolve().parent
w=openpyxl.load_workbook(p/'Dashboard_vs_Service_Confirmations.xlsx',data_only=True)
f=openpyxl.load_workbook(p/'Dashboard_vs_Service_Confirmations.xlsx',data_only=False)
d=json.loads((p/'comparison_data.json').read_text(encoding='utf-8'))
for s in w:
 for row in s:
  for c in row:
   assert c.data_type!='e',(s.title,c.coordinate,c.value)
   assert not(isinstance(c.value,str) and re.search(r'[\u4e00-\u9fff]',c.value)),(s.title,c.coordinate)
for name,count in [('Dashboard Tickets',3564),('Matched Tickets',76),('Report Details',473),('Invoice Lines',335)]:
 assert w[name].max_row-4==count
 assert len(f[name].tables)==1
 assert f[name].freeze_panes=='C5'
assert abs(w['Summary']['D10'].value-111.94)<1e-7
assert abs(w['Summary']['H10'].value-33205.79)<1e-7
assert abs(w['Summary']['D19'].value+99.67)<1e-7
assert abs(w['Summary']['H19'].value-11647.59)<1e-7
assert w['Summary']['B25'].value==76
assert abs(w['Summary']['C25'].value-265.25)<1e-7
assert abs(w['Summary']['H25'].value-1365.50)<1e-7
rows=list(w['Matched Tickets'].values)[4:]
assert all(abs(r[11])<1e-7 for r in rows)
diff=[r for r in rows if abs(r[14])>1e-7]
assert len(diff)==1 and diff[0][0]=='36147'
assert sum(r[15]=='Yes' for r in rows)==48
assert sum(r[16]=='Yes' for r in rows)==5
assert diff[0][5]=='0090043439'
assert diff[0][17]==0
print('Verified: dashboard monthly totals, report totals, 76 matched tickets, 1 invoice reversal, formulas, row counts and identifiers.')
