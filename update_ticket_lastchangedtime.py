from __future__ import annotations

import base64
import json
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
import os
import re
import ssl
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


WORKBOOK_PATH = Path(
    os.getenv(
        "SOURCE_WORKBOOK",
        str(Path(__file__).resolve().parent / "c4c_ticket_table_z007_z010_checked_hana_final.xlsx"),
    )
)
TARGET_SHEETS = ("Tickets", "NotAssigned")
DEPRECATED_COLUMNS = {"ChangeOnDateTime"}
TARGET_COLUMN = "lastchangedtime"
BASE_URL = os.getenv(
    "C4C_BASE_URL",
    "https://longcui-automobile-cpi-tyrbc1k7.it-cpi010-rt.cpi.cn40.apps.platform.sapcloud.cn",
)
HISTORY_PATH = os.getenv("C4C_HISTORY_API_PATH", "/http/PC4C/Ticket/getChangeHistory")
POSTMAN_COLLECTION = Path(
    os.getenv(
        "C4C_HISTORY_POSTMAN_COLLECTION",
        str(Path.home() / "Downloads" / "c4c-histroy.json"),
    )
)
POSTMAN_ENV_KEY = os.getenv("C4C_HISTORY_ENV", "PC4C").upper()
USERNAME = os.getenv("C4C_USERNAME", "")
PASSWORD = os.getenv("C4C_PASSWORD", "")
TIMEOUT = int(os.getenv("C4C_TIMEOUT", "30"))
VERIFY_SSL = os.getenv("C4C_VERIFY_SSL", "false").lower() in {"1", "true", "yes", "y"}
HISTORY_PAGE_SIZE = int(os.getenv("C4C_HISTORY_PAGE_SIZE", "500"))
HISTORY_MAX_PAGES = int(os.getenv("C4C_HISTORY_MAX_PAGES", "20"))
HISTORY_WORKERS = int(os.getenv("C4C_HISTORY_WORKERS", "16"))
FETCH_HISTORY = os.getenv("C4C_FETCH_HISTORY", "true").lower() in {"1", "true", "yes", "y"}
HISTORY_SCOPE = os.getenv("C4C_HISTORY_SCOPE", "dashboard").strip().lower()
FORCE_REFRESH = os.getenv("C4C_FORCE_HISTORY_REFRESH", "false").lower() in {"1", "true", "yes", "y"}
STATUS_CHANGE_FIELD = "Status/ServiceRequestLifeCycleStatusCode"
C4C_DATE_RE = re.compile(r"/Date\((-?\d+)\)/")
ROOT_NODE_RE = re.compile(r"^Root\([^)]+\)$")


def iter_postman_requests(items: list[dict[str, Any]]):
    stack = list(items)
    while stack:
        item = stack.pop(0)
        children = item.get("item")
        if isinstance(children, list):
            stack[:0] = children
        else:
            yield item


def postman_basic_auth_for_path(path_fragment: str) -> tuple[str, str] | None:
    if not POSTMAN_COLLECTION.exists():
        return None
    try:
        payload = json.loads(POSTMAN_COLLECTION.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None

    for item in iter_postman_requests(payload.get("item", [])):
        request = item.get("request", {})
        raw_url = str(request.get("url", {}).get("raw", ""))
        if path_fragment not in raw_url:
            continue
        auth = request.get("auth", {}) or item.get("auth", {}) or {}
        basic = {
            entry.get("key"): entry.get("value")
            for entry in auth.get("basic", [])
            if isinstance(entry, dict)
        }
        username = clean_text(basic.get("username"))
        password = clean_text(basic.get("password"))
        if username and password:
            return username, password
    return None


def resolve_history_auth() -> tuple[str, str]:
    if USERNAME and PASSWORD:
        return USERNAME, PASSWORD

    env_fragment = f"/http/{POSTMAN_ENV_KEY}/Ticket/getChangeHistory"
    found = postman_basic_auth_for_path(env_fragment)
    if found:
        return found

    found = postman_basic_auth_for_path(HISTORY_PATH)
    if found:
        return found

    raise SystemExit(
        "C4C history credentials not found. Set C4C_USERNAME/C4C_PASSWORD or "
        f"provide {POSTMAN_COLLECTION}."
    )


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() == "nan":
        return ""
    return text


def build_history_url(ticket_id: str, skip: int) -> str:
    query = urlencode({"srId": ticket_id, "pageSize": HISTORY_PAGE_SIZE, "skip": skip})
    return f"{BASE_URL.rstrip('/')}{HISTORY_PATH}?{query}"


def auth_header() -> str:
    username, password = resolve_history_auth()
    token = base64.b64encode(f"{username}:{password}".encode("utf-8")).decode("ascii")
    return f"Basic {token}"


def ssl_context() -> ssl.SSLContext | None:
    if VERIFY_SSL:
        return None
    return ssl._create_unverified_context()


def fetch_ticket_history_page(ticket_id: str, skip: int) -> list[dict[str, Any]]:
    request = Request(
        build_history_url(ticket_id, skip),
        headers={"Accept": "application/json", "Authorization": auth_header()},
        method="GET",
    )
    with urlopen(request, timeout=TIMEOUT, context=ssl_context()) as response:
        body = response.read().decode("utf-8")
    payload = json.loads(body)
    return list(payload.get("d", {}).get("results", []))


def fetch_ticket_history(ticket_id: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    skip = 0
    for _ in range(HISTORY_MAX_PAGES):
        page = fetch_ticket_history_page(ticket_id, skip)
        rows.extend(page)
        if len(page) < HISTORY_PAGE_SIZE:
            break
        skip += HISTORY_PAGE_SIZE
    return rows


def fetch_ticket_history_task(ticket_id: str) -> tuple[str, list[dict[str, Any]], str | None]:
    try:
        return ticket_id, fetch_ticket_history(ticket_id), None
    except Exception as exc:
        return ticket_id, [], str(exc)


def parse_c4c_change_datetime(value: Any) -> datetime | None:
    match = C4C_DATE_RE.search(clean_text(value))
    if not match:
        return None
    try:
        milliseconds = int(match.group(1))
    except ValueError:
        return None
    return datetime.fromtimestamp(milliseconds / 1000, tz=timezone.utc).replace(tzinfo=None)


def is_root_lifecycle_status_update(record: dict[str, Any]) -> bool:
    return (
        clean_text(record.get("ObjectNodeElementName")) == STATUS_CHANGE_FIELD
        and clean_text(record.get("ObjectNodeElementModificationTypeCode")).lower() == "update"
        and bool(ROOT_NODE_RE.match(clean_text(record.get("CompleteNodeHierarchy"))))
    )


def last_lifecycle_status_change(history_rows: list[dict[str, Any]]) -> str:
    status_changes = [
        value
        for value in (
            parse_c4c_change_datetime(row.get("ChangeDateTime"))
            for row in history_rows
            if is_root_lifecycle_status_update(row)
        )
        if value is not None
    ]
    if not status_changes:
        return ""
    return max(status_changes).strftime("%Y-%m-%d %H:%M:%S")


def header_map(ws) -> dict[str, int]:
    return {
        clean_text(cell.value): cell.column
        for cell in ws[1]
        if clean_text(cell.value)
    }


def copy_column_style(ws, source_col: int, target_col: int) -> None:
    for row in range(1, ws.max_row + 1):
        source = ws.cell(row=row, column=source_col)
        target = ws.cell(row=row, column=target_col)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.fill = copy(source.fill)
        target.font = copy(source.font)
        target.border = copy(source.border)


def delete_deprecated_columns(ws) -> list[str]:
    deleted = []
    headers = [clean_text(cell.value) for cell in ws[1]]
    for idx in range(len(headers), 0, -1):
        header = headers[idx - 1]
        if header in DEPRECATED_COLUMNS:
            ws.delete_cols(idx)
            deleted.append(header)
    return deleted


def ensure_lastchangedtime_column(ws) -> int:
    headers = header_map(ws)
    if TARGET_COLUMN in headers:
        return headers[TARGET_COLUMN]

    insert_at = headers.get("CreatedOn", ws.max_column) + 1
    ws.insert_cols(insert_at)
    copy_column_style(ws, max(1, insert_at - 1), insert_at)
    ws.cell(row=1, column=insert_at).value = TARGET_COLUMN
    ws.column_dimensions[ws.cell(row=1, column=insert_at).column_letter].width = 20
    return insert_at


def should_refresh_row(ws, row: int, headers: dict[str, int]) -> bool:
    if HISTORY_SCOPE == "all":
        return True

    def cell_text(header: str) -> str:
        col = headers.get(header)
        return clean_text(ws.cell(row=row, column=col).value) if col else ""

    ticket_type = cell_text("TicketType").upper()
    ticket_type_text = cell_text("TicketTypeText").lower()
    status = cell_text("TicketStatusText").lower()
    billing_date = cell_text("Billing date")
    is_pdi = ticket_type == "Z010" or ticket_type_text == "pdi"
    is_repair = ticket_type == "Z007" or ticket_type_text == "repair ticket"

    if is_pdi and status in {"repair completed", "claim time ticket"}:
        return True
    if is_repair and status in {"repair completed", "claim time ticket"}:
        return True
    if is_repair and status == "create invoice" and not billing_date:
        return True
    return False


def collect_ticket_ids(wb) -> list[str]:
    ticket_ids: list[str] = []
    seen = set()
    for sheet_name in TARGET_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = header_map(ws)
        ticket_col = headers.get("TicketID")
        if not ticket_col:
            continue
        lastchanged_col = headers.get(TARGET_COLUMN)
        for row in range(2, ws.max_row + 1):
            ticket_id = clean_text(ws.cell(row=row, column=ticket_col).value)
            existing = clean_text(ws.cell(row=row, column=lastchanged_col).value) if lastchanged_col else ""
            if existing and not FORCE_REFRESH:
                continue
            if ticket_id and ticket_id not in seen and should_refresh_row(ws, row, headers):
                seen.add(ticket_id)
                ticket_ids.append(ticket_id)
    return ticket_ids


def fetch_lastchanged_by_ticket(ticket_ids: list[str]) -> dict[str, str]:
    if not FETCH_HISTORY or not ticket_ids:
        return {}

    started = time.time()
    print(f"[RUN] Fetching status history for {len(ticket_ids)} tickets (scope={HISTORY_SCOPE})...", flush=True)
    lastchanged_by_ticket: dict[str, str] = {}
    failures: list[tuple[str, str]] = []
    workers = min(max(1, HISTORY_WORKERS), max(1, len(ticket_ids)))

    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(fetch_ticket_history_task, ticket_id): ticket_id
            for ticket_id in ticket_ids
        }
        for idx, future in enumerate(as_completed(futures), start=1):
            ticket_id, history_rows, error = future.result()
            if error:
                failures.append((ticket_id, error))
                continue
            lastchanged_by_ticket[ticket_id] = last_lifecycle_status_change(history_rows)
            if idx % 100 == 0 or idx == len(ticket_ids):
                print(f"[INFO] History progress: {idx}/{len(ticket_ids)}", flush=True)

    populated = sum(1 for value in lastchanged_by_ticket.values() if value)
    print(
        f"[OK] History complete in {time.time() - started:.1f}s; "
        f"populated={populated}, blank={len(ticket_ids) - populated}, failures={len(failures)}"
        ,
        flush=True,
    )
    if failures:
        print("[WARN] First history failures:", flush=True)
        for ticket_id, error in failures[:10]:
            print(f"       {ticket_id}: {error}", flush=True)
    if len(failures) == len(ticket_ids):
        raise SystemExit("History update aborted: all C4C history requests failed.")
    return lastchanged_by_ticket


def write_lastchangedtime(wb, lastchanged_by_ticket: dict[str, str]) -> dict[str, Any]:
    stats = {"deletedColumns": [], "targetSheets": 0, "updatedCells": 0}
    for sheet_name in TARGET_SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        deleted = delete_deprecated_columns(ws)
        if deleted:
            stats["deletedColumns"].append({sheet_name: deleted})

        lastchanged_col = ensure_lastchangedtime_column(ws)
        headers = header_map(ws)
        ticket_col = headers.get("TicketID")
        if not ticket_col:
            continue

        stats["targetSheets"] += 1
        for row in range(2, ws.max_row + 1):
            ticket_id = clean_text(ws.cell(row=row, column=ticket_col).value)
            if ticket_id not in lastchanged_by_ticket:
                continue
            value = lastchanged_by_ticket[ticket_id]
            ws.cell(row=row, column=lastchanged_col).value = value
            if value:
                stats["updatedCells"] += 1
    return stats


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise SystemExit(f"Source workbook not found: {WORKBOOK_PATH}")

    wb = load_workbook(WORKBOOK_PATH)
    ticket_ids = collect_ticket_ids(wb)
    schema_stats = write_lastchangedtime(wb, {})
    wb.save(WORKBOOK_PATH)

    lastchanged_by_ticket = fetch_lastchanged_by_ticket(ticket_ids)
    stats = write_lastchangedtime(wb, lastchanged_by_ticket)
    wb.save(WORKBOOK_PATH)

    print(
        "[OK] Workbook updated: "
        f"{WORKBOOK_PATH} "
        f"(sheets={stats['targetSheets'] or schema_stats['targetSheets']}, "
        f"updatedCells={stats['updatedCells']})"
    )


if __name__ == "__main__":
    main()
